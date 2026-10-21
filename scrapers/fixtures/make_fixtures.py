#!/usr/bin/env python3
"""Build every file in fixtures/ from the hand-copied samples below.

Run from anywhere:  python fixtures/make_fixtures.py
Re-running rewrites byte-identical files (XLSX zip timestamps and PDF metadata are pinned).

PROVENANCE KEY (also written to fixtures/provenance.json)
  REAL-VERBATIM   copied field-by-field from the live source via WebFetch on 2026-10-20
  REAL-VIA-GRIDSTATUS  real rows as rendered in gridstatus' documentation (a live run of
                  gridstatus against the source); raw key names from gridstatus source
                  code, commit 8a4d099 (2026-10-16). Fields the docs hid are left blank.
  SYNTHETIC       invented values in the verified column layout; used only where no real
                  rows were visible from this sandbox. Every synthetic name says SYNTHETIC.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
from pathlib import Path

HERE = Path(__file__).resolve().parent
PINNED = dt.datetime(2026, 10, 20, 0, 0, 0)

PROVENANCE: dict[str, dict] = {}


def note(path: str, kind: str, source: str, detail: str) -> None:
    PROVENANCE[path] = {"provenance": kind, "source": source, "detail": detail}


def save_xlsx(wb, path: Path) -> None:
    """Save an openpyxl workbook with pinned zip-entry timestamps so rebuilds are byte-identical."""
    import zipfile

    wb.properties.created = PINNED
    wb.properties.modified = PINNED
    raw = io.BytesIO()
    wb.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as zin, \
            zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=PINNED.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":  # openpyxl stamps 'modified' with the save time
                import re

                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>2026-10-20T00:00:00Z\g<2>", data)
            zout.writestr(info, data)
    path.write_bytes(out.getvalue())


# ======================================================================================
# ERCOT: render the verbatim April-2026 excerpt as a PDF (tests the pypdf path)
# ======================================================================================

def build_ercot() -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate

    txt = (HERE / "ercot" / "ERCOT-Monthly-April-2026-FINAL_p4_excerpt.txt").read_text().strip().split("\n")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, title="ERCOT Monthly April 2026 p4 excerpt (fixture)",
                            invariant=1)
    styles = getSampleStyleSheet()
    doc.build([Paragraph(txt[0], styles["Heading2"])] + [Paragraph(t, styles["BodyText"]) for t in txt[1:]])
    (HERE / "ercot" / "ERCOT-Monthly-April-2026-FINAL_p4_excerpt.pdf").write_bytes(buf.getvalue())
    for f, url in [
        ("ercot/ERCOT-Monthly-April-2026-FINAL_p4_excerpt.txt",
         "https://www.ercot.com/files/docs/2026/05/13/ERCOT-Monthly-April-2026-FINAL.pdf"),
        ("ercot/9-Interconnection-and-Grid-Analysis-Update_2026-04-13_slides2-3_excerpt.txt",
         "https://www.ercot.com/files/docs/2026/04/13/9-Interconnection-and-Grid-Analysis-Update.pdf"),
        ("ercot/ERCOT-Monthly-Operational-Overview-August-2026_p9_excerpt.txt",
         "https://www.ercot.com/files/docs/2026/09/16/ERCOT-Monthly-Operational-Overview-August-2026.pdf"),
    ]:
        note(f, "REAL-VERBATIM", url, "Large-load passage only; quoted fragments re-joined with single spaces.")
    note("ercot/ERCOT-Monthly-April-2026-FINAL_p4_excerpt.pdf", "REAL-VERBATIM",
         "https://www.ercot.com/files/docs/2026/05/13/ERCOT-Monthly-April-2026-FINAL.pdf",
         "The verbatim .txt excerpt typeset into a one-page PDF with reportlab (layout is not ERCOT's).")


# ======================================================================================
# SPP: GenerateSummaryCSV, 16 rows copied field-by-field (REAL-VERBATIM)
# ======================================================================================

SPP_HEADER = [
    "Generation Interconnection Number", "IFS Queue Number", "Current Cluster", "Cluster Group",
    # gridstatus' parser (spp.py) renames " Nearest Town or County" WITH a leading space, and its
    # docs show County populated, so the raw header carries that space. WebFetch trims it.
    " Nearest Town or County",
    "State", "TO at POI", "In-Service Date", "Commercial Operation Date", "Cessation Date",
    "Original Generator Commercial Op Date", "Capacity", "MAX Summer MW", "MAX Winter MW", "Service Type",
    "Requested Maximum Injection Capability (MW)", "Requested Network Resource Deliverability (MW)",
    "Nameplate Capacity", "Generation Type", "Fuel Type", "Substation or Line", "Request Received",
    "Date Withdrawn", "Status", "JTIQ Participant", "JTIQ Commitment", "Cause of Delay",
]
_D = "GIA delayed due to  Affected Systems Study"
# (id, cluster, group, county, state, to, isd, cod, cess, cap, summ, wint, svc, rmic, rnrd, nameplate,
#  gentype, fuel, poi, received, withdrawn, status, jtiq, delay)
SPP_ROWS = [
    ("TI-18-0827", "RTOE Transitional Cluster", "", "Weld County", "CO", "TSGT", "6/24/2022", "", "", "145", "145", "145", "ER/NR", "0", "0", "0", "Wind", "Wind", "Redtail 115  Substation", "8/27/2018", "", "IA FULLY EXECUTED/COMMERCIAL OPERATION", "", ""),
    ("TI-17-0224", "RTOE Transitional Cluster", "", "Kit Carson/Cheyenne Counties", "CO", "TSGT", "9/2/2020", "", "", "104", "104", "104", "ER/NR", "0", "0", "0", "Wind", "Wind", "Windtalker 230 Switching Station", "2/24/2017", "", "IA FULLY EXECUTED/COMMERCIAL OPERATION", "", ""),
    ("TI-12-0217", "RTOE Transitional Cluster", "", "Ouray County", "CO", "TSGT", "11/1/2013", "", "", "7.2", "7.2", "7.2", "ER/NR", "0", "0", "0", "Solar", "Photovoltaic", "Cow Creek 115 Switching Station", "2/17/2012", "", "IA FULLY EXECUTED/COMMERCIAL OPERATION", "", ""),
    ("TI-08-0502", "RTOE Transitional Cluster", "", "Kit Carson County", "CO", "TSGT", "10/15/2010", "", "", "51", "51", "51", "ER/NR", "0", "0", "0", "Wind", "Wind", "Landsman Creek 230 SwitchingStation", "5/2/2008", "", "IA FULLY EXECUTED/COMMERCIAL OPERATION", "", ""),
    ("TI-08-0204", "RTOE Transitional Cluster", "", "Kit Carson County", "CO", "TSGT", "11/30/2015", "", "", "150", "150", "150", "ER/NR", "0", "0", "0", "Wind", "Wind", "Burlington 230  Bus", "2/4/2008", "", "IA FULLY EXECUTED/COMMERCIAL OPERATION", "", ""),
    ("GI-TC-2024-31", "RTOE Transitional Cluster", "", "Elbert County", "CO", "TSGT", "10/1/2027", "12/31/2028", "", "398", "398", "398", "ER", "0", "0", "0", "Hybrid", "Solar/Storage", "Lincoln - Midway 230  Line", "5/29/2024", "", "IA FULLY EXECUTED/ON SCHEDULE", "", _D),
    ("GI-TC-2024-30", "RTOE Transitional Cluster", "", "Morgan County", "CO", "TSGT", "3/1/2028", "", "", "199.9", "199.9", "199.9", "ER/NR", "0", "0", "0", "Hybrid", "Solar/Storage", "Story 230 Substation", "5/29/2024", "", "IA FULLY EXECUTED/ON SCHEDULE", "", _D),
    ("GI-TC-2024-29", "RTOE Transitional Cluster", "", "Kit Carson County", "CO", "TSGT", "9/15/2027", "10/1/2028", "", "170", "170", "170", "ER/NR", "0", "0", "0", "Wind", "Wind", "Landsman Creek-Windtalker 230Line", "5/29/2024", "", "FACILITY STUDY STAGE", "", _D),
    ("GI-TC-2024-28", "RTOE Transitional Cluster", "", "Weld County", "CO", "TSGT", "6/2/2031", "12/31/2027", "", "222.3", "222.3", "222.3", "ER/NR", "0", "0", "0", "Solar", "Photovoltaic", "Lloyd Sisson 115  Substation", "5/31/2023", "", "FACILITY STUDY STAGE", "", _D),
    ("GI-TC-2024-26", "RTOE Transitional Cluster", "", "Elbert County", "CO", "TSGT", "1/1/2028", "8/17/2027", "", "199", "199", "199", "ER/NR", "0", "0", "0", "Hybrid", "Solar/Storage", "Lincoln – Midway 230  Line", "5/31/2023", "", "IA FULLY EXECUTED/ON SCHEDULE", "", _D),
    ("GI-TC-2024-25", "RTOE Transitional Cluster", "", "Weld County", "CO", "TSGT", "6/2/2031", "12/1/2028", "", "350", "350", "350", "ER", "0", "0", "0", "Wind", "Wind", "Ault - Laramie River Station 345\nLine", "5/28/2024", "", "FACILITY STUDY STAGE", "", _D),
    ("GI-TC-2024-22", "RTOE Transitional Cluster", "", "Weld County", "CO", "TSGT", "6/2/2031", "3/1/2028", "", "140", "140", "140", "ER/NR", "0", "0", "0", "Solar", "Photovoltaic", "Lloyd Sisson 115  Substation", "5/23/2024", "", "FACILITY STUDY STAGE", "", _D),
    ("GI-CS-2025-27", "RTOE Transitional Cluster", "", "Delta County,", "CO", "TSGT", "", "4/1/2029", "", "199", "199", "199", "ER/NR", "0", "0", "0", "Thermal", "Gas", "Starr Nelson 115 Substation", "10/3/2025", "", "WITHDRAWN", "", ""),
    ("GEN-2024-GR1", "Replacement", "04 SOUTHEAST", "Harrah", "OK", "OG&E", "", "12/31/2026", "12/31/2026", "492", "492", "492", "ER", "0", "0", "0", "Thermal", "CTG", "Horseshoe Lake 138kV", "1/8/2024", "", "IA FULLY EXECUTED/ON SCHEDULE", "N/A", ""),
    ("GEN-2024-379", "DISIS-2024-001", "04 SOUTHEAST", "Bossier Parish", "LA", "AEP", "5/1/2030", "12/15/2031", "", "1350", "1310", "1350", "ER/NR", "1350", "1350", "1694.135", "Thermal", "Gas", "Longwood - El Dorado 345 kV Line", "3/1/2025", "", "DISIS STAGE", "Yes", ""),
    ("GEN-2024-377", "DISIS-2024-001", "04 SOUTHEAST", "Harrison", "TX", "AEP", "1/5/2030", "12/15/2031", "", "1350", "1310", "1350", "ER/NR", "1350", "1350", "1694.135", "Thermal", "Gas", "Pirkey 345 kV Substation", "3/1/2025", "", "DISIS STAGE", "Yes", ""),
]


def build_spp() -> None:
    buf = io.StringIO()
    buf.write('"Last Updated On",10/19/2026,\r\n')  # exactly as WebFetch quoted the first line
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(SPP_HEADER)
    for r in SPP_ROWS:
        (qid, cluster, group, county, state, to, isd, cod, cess, cap, summ, wint, svc, rmic, rnrd, npc,
         gtype, fuel, poi, recv, wd, status, jtiq, delay) = r
        w.writerow([qid, "", cluster, group, county, state, to, isd, cod, cess, "", cap, summ, wint, svc,
                    rmic, rnrd, npc, gtype, fuel, poi, recv, wd, status, jtiq, "", delay])
    (HERE / "spp" / "spp_GenerateSummaryCSV_sample.csv").write_text(buf.getvalue(), encoding="utf-8")
    note("spp/spp_GenerateSummaryCSV_sample.csv", "REAL-VERBATIM",
         "https://opsportal.spp.org/Studies/GenerateSummaryCSV",
         "Title line + 27-column header + first 12 data rows + first 4 Thermal (Gas/CTG) rows, file dated "
         "10/19/2026. Field values copied one by one; CSV quoting style of the real file not observed.")


# ======================================================================================
# MISO: getprojects JSON (REAL-VIA-GRIDSTATUS)
# ======================================================================================

MISO_KEYS = [  # raw keys used by gridstatus/miso.py (commit 8a4d099, 2026-10-16)
    "projectNumber", "county", "state", "transmissionOwner", "poiName", "queueDate", "withdrawnDate",
    "applicationStatus", "summerNetMW", "winterNetMW", "negInService", "fuelType", "facilityType",
    "postGIAStatus", "doneDate", "inService", "giaToExec", "studyCycle", "studyGroup", "studyPhase",
    "svcType", "dp1ErisMw", "dp1NrisMw", "dp2ErisMw", "dp2NrisMw", "sisPhase1",
]


def _miso(**kw):
    rec = {k: None for k in MISO_KEYS}
    rec.update({"dp1ErisMw": 0.0, "dp1NrisMw": 0.0, "dp2ErisMw": 0.0, "dp2NrisMw": 0.0})
    rec.update(kw)
    return rec


def build_miso() -> None:
    blank = dict(county="", state="", poiName="", fuelType="", studyPhase="", sisPhase1="")
    rows = [
        # latest gridstatus docs build (2026): first and last rows of MISO table
        _miso(projectNumber="J4182", transmissionOwner="ENTERGY ARKANSAS, LLC", summerNetMW=0.0, winterNetMW=0.0,
              studyCycle="ERAS Study Cycle 2", studyGroup="South", svcType="NRIS", **blank),
        _miso(projectNumber="J4176", transmissionOwner="ENTERGY MISSISSIPPI, LLC.", summerNetMW=1750.0,
              studyCycle="DPP-2025", studyGroup="South", svcType="NRIS", **{**blank, "studyPhase": "Phase 1"}),
        _miso(projectNumber="J4179", transmissionOwner="AMERICAN TRANSMISSION COMPANY", summerNetMW=735.0,
              studyCycle="DPP-2026", studyGroup="East (ATC)", svcType="NRIS", **blank),
        _miso(projectNumber="J4184", transmissionOwner="AMERICAN TRANSMISSION COMPANY", summerNetMW=755.0,
              studyCycle="DPP-2026", studyGroup="East (ATC)", svcType="NRIS", **blank),
        _miso(projectNumber="J4183", transmissionOwner="ENTERGY MISSISSIPPI, LLC.", summerNetMW=99.9,
              studyCycle="DPP-2025", studyGroup="South", svcType="NRIS", dp1ErisMw=23.0, dp1NrisMw=23.0,
              **{**blank, "studyPhase": "Phase 1"}),
        _miso(projectNumber="S1067", county="Washtenaw", state="MI", poiName="Majestic- Lemoyne 345 kV",
              transmissionOwner="ITC", summerNetMW=75.0, svcType="SIS", fuelType="", studyCycle="", studyGroup="",
              studyPhase="", sisPhase1=""),
        _miso(projectNumber="S1046", county="Codington", state="SD", poiName="Big Stone South 230kV",
              transmissionOwner="OTTER TAIL POWER COMPANY", summerNetMW=120.0, svcType="SIS", fuelType="",
              studyCycle="", studyGroup="", studyPhase="", sisPhase1=""),
        _miso(projectNumber="S1045", county="Emmons", state="ND", poiName="Napoleon SW 230 kV (661091)",
              transmissionOwner="MONTANA-DAKOTA UTILITIES CO.", summerNetMW=170.0, svcType="SIS", fuelType="",
              studyCycle="", studyGroup="", studyPhase="", sisPhase1=""),
        _miso(projectNumber="S1054", county="Cottonwood", state="MN", poiName="Lakefield Junction 345kV",
              transmissionOwner="NORTHERN STATES POWER COMPANY", summerNetMW=300.0, svcType="SIS", fuelType="",
              studyCycle="", studyGroup="", studyPhase="", sisPhase1=""),
        _miso(projectNumber="R5062", county="Jasper", state="IN", poiName="SCHAHFER GEN STA",
              transmissionOwner="NORTHERN INDIANA PUBLIC SERVICE COMPANY LLC", summerNetMW=732.73,
              fuelType="Hybrid", svcType="", studyCycle="", studyGroup="", studyPhase="", sisPhase1=""),
        # older gridstatus docs (0.13.0 all-ISO table) - raw date strings as rendered there
        _miso(projectNumber="E002", county="Pointe Coupee", state="LA", poiName="500 kV switchyard at Big Cajon II",
              transmissionOwner="Entergy", fuelType="Gas", summerNetMW=15.0, winterNetMW=15.0,
              queueDate="2013-09-13T04:00:00Z", applicationStatus="Done"),
        _miso(projectNumber="E291", county="", state="TX", poiName="Rich 138 kV substation",
              transmissionOwner="Entergy", fuelType="Hydro", summerNetMW=27.0, winterNetMW=27.0,
              queueDate="2012-05-15T04:00:00Z", applicationStatus="Done", negInService="2016-01-31T05:00:00Z"),
        _miso(projectNumber="G001", county="Brown County", state="WI", poiName="Unknown",
              transmissionOwner="American Transmission Co. LLC", fuelType="", summerNetMW=255.0, winterNetMW=255.0,
              queueDate="1995-11-08T05:00:00Z", applicationStatus="Withdrawn", negInService="2014-11-17T05:00:00Z"),
    ]
    (HERE / "miso" / "miso_getprojects_sample.json").write_text(json.dumps(rows, indent=1) + "\n")
    note("miso/miso_getprojects_sample.json", "REAL-VIA-GRIDSTATUS",
         "https://www.misoenergy.org/api/giqueue/getprojects (WebFetch got a client error; not seen directly)",
         "Values from gridstatus docs tables (latest build: J4182..R5062; 0.13.0 build: E002, E291, G001). "
         "Raw key names from gridstatus/miso.py. applicationStatus/queueDate for the J/S/R rows were hidden "
         "in the docs and are null here. Exact JSON value types (string vs number) not observed.")


# ======================================================================================
# NYISO: NYISO-Interconnection-Queue.xlsx (REAL-VIA-GRIDSTATUS values, gridstatus raw names)
# ======================================================================================

NYISO_ACTIVE_HEADER = [  # names from gridstatus/nyiso.py; ORDER NOT VERIFIED for these sheets
    "Queue Pos.", "Project Name", "Date of IR", "SP (MW)", "WP (MW)", "Type/ Fuel", "County", "State", "Z",
    "Points of Interconnection", "Utility", "S", "Last Updated Date", "Availability of Studies",
    "SGIA Tender Date", "Proposed In-Service Date", "Proposed Initial-Sync Date", "Proposed COD",
]
NYISO_WITHDRAWN_HEADER = [
    "Queue Pos.", "Owner/Developer", "Project Name", "Date of IR", "SP (MW)", "WP (MW)", "Type/ Fuel",
    "County", "State", "Z", "Points of Interconnection", "Utility ", "S", "Last Update", "Availability of Studies",
]
# In Service sheet: two-row merged header. The first 15 columns' order IS implied by gridstatus
# (it inserts new columns at positions 15-17 and names ("Z", "Unnamed: 9_level_1") etc.).
NYISO_INSERVICE_TOP = ["Queue", None, None, "Date", "SP", "WP", "Type/", "Location", None, "Z",
                       "Interconnection", None, None, "Last Update", "Availability", "Proposed", None, "Status"]
NYISO_INSERVICE_SUB = ["Pos.", "Owner/Developer", "Project Name", "of IR", "(MW)", "(MW)", "Fuel", "County",
                       "State", None, "Point", "Utility ", "S", None, "of Studies", " In-Service", "COD", None]


def build_nyiso() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Interconnection Queue"
    ws.append(NYISO_ACTIVE_HEADER)
    active = [  # (pos, name, sp, fuel_code, county, poi, utility, last_upd, studies)
        ("0276", "Homer Solar Energy Center", 90.0, "S", "Cortland", "Cortland - Fenner 115kV", "NM-NG", "2026-05-31", "FES, SRIS, FS"),
        ("0495", "Mohawk Solar", 90.5, "S", "Montgomery", "St. Johnsville - Marshville 115kV", "NM-NG", "2025-09-30", "FES, SRIS, FS"),
        ("0519", "Canisteo Wind", 290.7, "W", "Steuben", "Bennett 115kV", "NYSEG", "2025-03-31", "FES, SRIS, FS"),
        ("0522", "NYC Energy", 79.9, "ES", "Kings", "Hudson Avenue East 138kV", "ConEd", "2026-06-30", "FES, SRIS, FS"),
        ("0535", "Riverhead Solar 2", 36.0, "S", "Suffolk", "Edwards 138kV", "LIPA", "2024-09-30", "FES, SRIS, FS"),
    ]
    for pos, name, sp, fuel, county, poi, util, upd, studies in active:
        ws.append([pos, name, None, sp, None, fuel, county, "NY", None, poi, util, None,
                   dt.datetime.fromisoformat(upd), studies, None, None, None, None])
    ws.append([None] * len(NYISO_ACTIVE_HEADER))
    ws.append(["Note: fixture footer row (real file has notes below the table)."])

    ws2 = wb.create_sheet(" Cluster Projects")  # leading space, as gridstatus reads it
    ws2.append(NYISO_ACTIVE_HEADER)
    ws3 = wb.create_sheet("Withdrawn")
    ws3.append(NYISO_WITHDRAWN_HEADER)
    ws4 = wb.create_sheet("Cluster Projects-Withdrawn")
    ws4.append(NYISO_WITHDRAWN_HEADER)

    ws5 = wb.create_sheet("In Service")
    ws5.append(NYISO_INSERVICE_TOP)
    ws5.append(NYISO_INSERVICE_SUB)
    for a, b in [("A1", "C1"), ("H1", "I1"), ("K1", "M1"), ("P1", "Q1")]:
        ws5.merge_cells(f"{a}:{b}")
    done = [  # (pos, dev, name, sp, fuel, county, z, poi, util, s, last_upd, studies)
        ("0769", "NY Power Authority", "North County Energy Storage", 20.0, "ES", "Franklin", "D", "Willis 115kV", "NYPA", 14, "2023-09-30", "SIS"),
        ("0775", "Puckett Solar, LLC", "Puckett Solar", 20.0, "S", "Chenango", "C", "Chenango Forks 34.5kV", "NYSEG", 14, "2022-09-30", "SIS, FS"),
        ("0781", "NY Power Authority", "Marcy South SSR Detection", 0.0, "AC", "Delaware", "E", "Fraser - Annex 345 kV Substation", "NYPA", 14, "2019-09-30", "SIS, FS"),
        ("0887", "CHPE LLC", "CH Uprate", 250.0, "DC", "Quebec-NY", "J", "Astoria Annex 345kV", "NYPA", 14, "2023-06-30", "SRIS, FS"),
    ]
    for pos, dev, name, sp, fuel, county, z, poi, util, s, upd, studies in done:
        ws5.append([pos, dev, name, None, sp, None, fuel, county, "NY", z, poi, util, s,
                    dt.datetime.fromisoformat(upd), studies, None, None, "In Service"])
    save_xlsx(wb, HERE / "nyiso" / "NYISO-Interconnection-Queue_sample.xlsx")
    note("nyiso/NYISO-Interconnection-Queue_sample.xlsx", "REAL-VIA-GRIDSTATUS",
         "https://www.nyiso.com/documents/20142/1407078/NYISO-Interconnection-Queue.xlsx (binary; not readable via WebFetch)",
         "Sheet names and column names from gridstatus/nyiso.py; values from gridstatus docs (2026 build), raw "
         "Type/ Fuel codes recovered by inverting gridstatus' code map (S, W, ES, AC, DC). Column ORDER on the "
         "active/withdrawn sheets is not verified. Cells the docs showed as NaN are blank. 'Status' text on the "
         "In Service sheet is a placeholder. No NYISO load (Type 'L') rows were visible, so none are included.")


# ======================================================================================
# PJM: queue export XLSX (REAL-VERBATIM values from the Cycle page table)
# ======================================================================================

PJM_HEADER = [  # export column names per gridstatus/pjm.py; order not verified
    "Project ID", "Name", "Commercial Name", "State", "County", "Status", "Transmission Owner", "MFO",
    "MW Energy", "MW Capacity", "MW In Service", "Fuel", "Submitted Date", "Projected In Service Date",
    "Actual In Service Date", "Withdrawal Date", "Withdrawn Remarks",
]
PJM_ROWS = [  # Project ID | Name | State | Status | TO | MFO | MW Energy | MW Capacity | MW In Service | Fuel
    ("AE1-070", "Elwood 345 kV", "IL", "Withdrawn", "ComEd", 1610.0, 0.0, 135.0, None, "Natural Gas"),
    ("AE1-092", "Blue Jacket-Kirby 138 kV", "OH", "EP", "Dayton", 206.55, 206.55, 96.4, None, "Solar"),
    ("AE1-114", "Maryland-Lancaster 138 kV", "IL", "EP", "ComEd", 150.0, 150.0, 34.0, None, "Wind"),
    ("AE1-148", "Kerr Dam-Ridge Rd 115 kV", "VA", "Withdrawn", "Dominion", 90.0, 90.0, 54.0, None, "Solar"),
    ("AE1-172", "Loretto-Wilton Center", "IL", "Withdrawn", "ComEd", 255.0, 255.0, 44.88, None, "Wind"),
    ("AE2-053", "Kerr Dam-Ridge Road 115 kV", "VA", "Withdrawn", "Dominion", 110.0, 20.0, 20.0, None, "Storage"),
    ("AE2-107", "Haumesser Road 138 kV", "IL", "Withdrawn", "ComEd", 90.0, 90.0, 54.0, None, "Solar"),
    ("AE2-123", "Birdneck-Landstown 230 kV", "VA", "Withdrawn", "Dominion", 800.1, 800.1, 163.3, None, "Offshore Wind"),
    ("AE2-124", "Landstown 230 kV", "VA", "Withdrawn", "Dominion", 800.1, 800.1, 148.5, None, "Offshore Wind"),
    ("AE2-130", "Rockport 765 kV", "IN", "Withdrawn", "AEP", 800.0, 800.0, 480.0, None, "Solar, Storage"),
    ("AE2-138", "Avon-North Clark 345 kV", "KY", "Withdrawn", "EKPC", 260.0, 260.0, 20.8, None, "Solar"),
    ("AE2-147", "Swamp 230 kV", "NC", "Withdrawn", "Dominion", 150.0, 150.0, 90.0, None, "Solar"),
    ("AE2-156", "Yadkin 115 kV", "VA", "EP", "Dominion", 100.0, 100.0, 100.0, None, "Storage"),
    ("AE2-173", "McLean 345 kV", "IL", "Withdrawn", "ComEd", 250.0, 50.0, 50.0, None, "Storage"),
    ("AE2-185", "Gladys DP-Stonemill Switching Station 69 kV", "VA", "EP", "Dominion", 60.0, 60.0, 36.0, None, "Solar"),
]


def build_pjm() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(PJM_HEADER)
    for pid, name, st, status, to, mfo, mwe, mwc, mwis, fuel in PJM_ROWS:
        ws.append([pid, name, None, st, None, status, to, mfo, mwe, mwc, mwis, fuel, None, None, None, None, None])
    save_xlsx(wb, HERE / "pjm" / "pjm_queue_export_sample.xlsx")
    note("pjm/pjm_queue_export_sample.xlsx", "REAL-VERBATIM",
         "https://www.pjm.com/planning/m/cycle-service-request-status (HTML table, first 15 rows)",
         "Values for Project ID, Name, State, Status, TO, MFO, MW Energy, MW Capacity, MW In Service, Fuel copied "
         "from the page table; written under the XLS-export column names used by gridstatus/pjm.py "
         "('TO' -> 'Transmission Owner'). Other export columns left blank (not seen). AE2-130 showed two fuel "
         "icons, written as 'Solar, Storage' (the export's own multi-fuel format was not seen).")


# ======================================================================================
# EIA: 860M XLSX and API v2 JSON (SYNTHETIC values, verified layouts)
# ======================================================================================

EIA_OPERATING = [
    "Entity ID", "Entity Name", "Plant ID", "Plant Name", "Plant State", "County", "Balancing Authority Code",
    "Sector", "Generator ID", "Unit Code", "Nameplate Capacity (MW)", "Net Summer Capacity (MW)",
    "Net Winter Capacity (MW)", "Technology", "Energy Source Code", "Prime Mover Code", "Operating Month",
    "Operating Year", "Planned Retirement Month", "Planned Retirement Year", "Status",
    "Nameplate Energy Capacity (MWh)", "DC Net Capacity (MW)", "Planned Derate Year", "Planned Derate Month",
    "Planned Derate of Summer Capacity (MW)", "Planned Uprate Year", "Planned Uprate Month",
    "Planned Uprate of Summer Capacity (MW)", "Latitude", "Longitude", "Google Map", "Bing Map",
]
EIA_PLANNED = [
    "Entity ID", "Entity Name", "Plant ID", "Plant Name", "Plant State", "County", "Balancing Authority Code",
    "Sector", "Generator ID", "Unit Code", "Nameplate Capacity (MW)", "Net Summer Capacity (MW)",
    "Net Winter Capacity (MW)", "Technology", "Energy Source Code", "Prime Mover Code",
    "Planned Operation Month", "Planned Operation Year", "Status", "Latitude", "Longitude",
]
EIA_RETIRED = [
    "Entity ID", "Entity Name", "Plant ID", "Plant Name", "Plant State", "County", "Balancing Authority Code",
    "Sector", "Generator ID", "Unit Code", "Nameplate Capacity (MW)", "Net Summer Capacity (MW)",
    "Net Winter Capacity (MW)", "Technology", "Energy Source Code", "Prime Mover Code", "Operating Month",
    "Operating Year", "Retirement Month", "Retirement Year", "Nameplate Energy Capacity (MWh)",
    "DC Net Capacity (MW)", "Latitude", "Longitude",
]
# (plant_id, plant, state, county, gen, nameplate, summer, tech, esc, pm, op_month, op_year)
EIA_OP_ROWS = [
    (990001, "SYNTHETIC GAS PLANT TX-A", "TX", "SYNTH", "CT1", 250.0, 230.0, "Natural Gas Fired Combustion Turbine", "NG", "GT", 6, 2023),
    (990001, "SYNTHETIC GAS PLANT TX-A", "TX", "SYNTH", "CT2", 250.0, 230.0, "Natural Gas Fired Combustion Turbine", "NG", "GT", 6, 2024),
    (990002, "SYNTHETIC SOLAR TX-B", "TX", "SYNTH", "PV1", 300.0, 300.0, "Solar Photovoltaic", "SUN", "PV", 3, 2024),
    (990003, "SYNTHETIC OLD GAS TX-C", "TX", "SYNTH", "ST1", 500.0, 480.0, "Natural Gas Steam Turbine", "NG", "ST", 1, 1985),
    (990004, "SYNTHETIC GAS PLANT VA-A", "VA", "SYNTH", "CC1", 600.0, 560.0, "Natural Gas Fired Combined Cycle", "NG", "CT", 12, 2022),
    (990005, "SYNTHETIC GAS PLANT GA-A", "GA", "SYNTH", "CT1", 180.0, 170.0, "Natural Gas Fired Combustion Turbine", "NG", "GT", 5, 2025),
    (990006, "SYNTHETIC GAS PLANT OH-A", "OH", "SYNTH", "CT1", 100.0, 95.0, "Natural Gas Internal Combustion Engine", "NG", "IC", 8, 2021),
    (990007, "SYNTHETIC GAS PLANT AZ-A", "AZ", "SYNTH", "CT1", 200.0, 190.0, "Natural Gas Fired Combustion Turbine", "NG", "GT", 7, 2026),
    (990008, "SYNTHETIC BATTERY PA-A", "PA", "SYNTH", "B1", 50.0, 50.0, "Batteries", "MWH", "BA", 9, 2024),
    (990009, "SYNTHETIC GAS PLANT WI-A", "WI", "SYNTH", "RICE1", 120.0, 118.0, "Natural Gas Internal Combustion Engine", "NG", "IC", 4, 2023),
    (990010, "SYNTHETIC GAS PLANT LA-A", "LA", "SYNTH", "CC1", 750.0, 700.0, "Natural Gas Fired Combined Cycle", "NG", "CT", 2, 2025),
    (990011, "SYNTHETIC GAS PLANT CA-A", "CA", "SYNTH", "CT1", 999.0, 990.0, "Natural Gas Fired Combustion Turbine", "NG", "GT", 1, 2024),
]
EIA_PLANNED_ROWS = [
    (990012, "SYNTHETIC PLANNED GAS TX-D", "TX", "SYNTH", "CC1", 1200.0, 1150.0, "Natural Gas Fired Combined Cycle", "NG", "CT", 6, 2028),
    (990013, "SYNTHETIC PLANNED GAS LA-B", "LA", "SYNTH", "CT1", 400.0, 380.0, "Natural Gas Fired Combustion Turbine", "NG", "GT", 1, 2027),
]
EIA_RETIRED_ROWS = [  # entered 2022, retired 2025: still counts as a 2022 addition
    (990014, "SYNTHETIC SHORT-LIVED GAS VA-B", "VA", "SYNTH", "IC1", 20.0, 19.0, "Natural Gas Internal Combustion Engine", "NG", "IC", 3, 2022, 6, 2025),
]


def build_eia() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    title = "U.S. Energy Information Administration - SYNTHETIC FIXTURE in EIA-860M layout (not real data)"
    ws = wb.active
    ws.title = "Operating"
    ws.append([title])
    ws.append(EIA_OPERATING)
    for pid, plant, st, county, gen, npc, summ, tech, esc, pm, om, oy in EIA_OP_ROWS:
        row = {c: None for c in EIA_OPERATING}
        row.update({"Entity ID": 99999, "Entity Name": "SYNTHETIC UTILITY", "Plant ID": pid, "Plant Name": plant,
                    "Plant State": st, "County": county, "Sector": "SYNTHETIC", "Generator ID": gen,
                    "Nameplate Capacity (MW)": npc, "Net Summer Capacity (MW)": summ, "Net Winter Capacity (MW)": summ,
                    "Technology": tech, "Energy Source Code": esc, "Prime Mover Code": pm, "Operating Month": om,
                    "Operating Year": oy, "Status": "(OP) Operating"})
        ws.append([row[c] for c in EIA_OPERATING])
    ws.append(["Note: SYNTHETIC footer row, mirroring the real file's trailing note."])

    ws = wb.create_sheet("Planned")
    ws.append([title])
    ws.append(EIA_PLANNED)
    for pid, plant, st, county, gen, npc, summ, tech, esc, pm, om, oy in EIA_PLANNED_ROWS:
        row = {c: None for c in EIA_PLANNED}
        row.update({"Entity ID": 99999, "Entity Name": "SYNTHETIC UTILITY", "Plant ID": pid, "Plant Name": plant,
                    "Plant State": st, "County": county, "Generator ID": gen, "Nameplate Capacity (MW)": npc,
                    "Net Summer Capacity (MW)": summ, "Technology": tech, "Energy Source Code": esc,
                    "Prime Mover Code": pm, "Planned Operation Month": om, "Planned Operation Year": oy,
                    "Status": "(U) Under construction"})
        ws.append([row[c] for c in EIA_PLANNED])
    ws.append(["Note: SYNTHETIC footer row."])

    ws = wb.create_sheet("Retired")
    ws.append([title])
    ws.append(EIA_RETIRED)
    for pid, plant, st, county, gen, npc, summ, tech, esc, pm, om, oy, rm, ry in EIA_RETIRED_ROWS:
        row = {c: None for c in EIA_RETIRED}
        row.update({"Entity ID": 99999, "Entity Name": "SYNTHETIC UTILITY", "Plant ID": pid, "Plant Name": plant,
                    "Plant State": st, "County": county, "Generator ID": gen, "Nameplate Capacity (MW)": npc,
                    "Net Summer Capacity (MW)": summ, "Technology": tech, "Energy Source Code": esc,
                    "Prime Mover Code": pm, "Operating Month": om, "Operating Year": oy,
                    "Retirement Month": rm, "Retirement Year": ry})
        ws.append([row[c] for c in EIA_RETIRED])
    ws.append(["Note: SYNTHETIC footer row."])
    wb.create_sheet("Canceled or Postponed").append([title])
    save_xlsx(wb, HERE / "eia" / "august_generator2026_SYNTHETIC.xlsx")
    note("eia/august_generator2026_SYNTHETIC.xlsx", "SYNTHETIC",
         "layout of https://www.eia.gov/electricity/data/eia860m/xls/august_generator2026.xlsx",
         "Sheet names (Operating/Planned/Retired/Canceled or Postponed), 1 title row, header row and trailing "
         "note row follow gridstatus/eia.py (skiprows=1, skipfooter=1); column names from gridstatus "
         "eia_constants.py with the '(MW)' suffixes it strips restored. Values invented.")

    rows = []
    for loc, desc in [("TX", "Texas"), ("VA", "Virginia")]:
        for period, all_gen, ng_gen in [("2026-06", 50000.0, 25000.0), ("2026-07", 56000.0, 30000.0)]:
            scale = 1.0 if loc == "TX" else 0.2
            for fid, fdesc, val in [("ALL", "all fuels", all_gen * scale), ("NG", "natural gas", ng_gen * scale)]:
                rows.append({"period": period, "location": loc, "stateDescription": desc, "sectorid": "99",
                             "sectorDescription": "All Sectors", "fueltypeid": fid, "fuelTypeDescription": fdesc,
                             "generation": round(val, 3), "generation-units": "thousand megawatthours"})
    payload = {
        "_fixture_note": "SYNTHETIC values in the EIA API v2 envelope (response.total/data). Row keys = facet ids "
                         "(location, sectorid, fueltypeid) + data id (generation) + '<id>-units', following the "
                         "convention visible in EIA's documented retail-sales example and gridstatus' facility-fuel "
                         "parser. Description keys are not used by the parser.",
        "response": {"total": str(len(rows)), "dateFormat": "YYYY-MM", "frequency": "monthly", "data": rows,
                     "description": "SYNTHETIC"},
        "request": {"command": "/v2/electricity/electric-power-operational-data/data/"},
        "apiVersion": "2.1.x",
    }
    (HERE / "eia" / "eia_v2_electric_power_operational_data_SYNTHETIC.json").write_text(json.dumps(payload, indent=1) + "\n")
    metadata = {
        "_fixture_note": "SYNTHETIC route-metadata response, used to test the facet/data-column check.",
        "response": {
            "id": "electric-power-operational-data",
            "facets": [{"id": "location"}, {"id": "sectorid"}, {"id": "fueltypeid"}],
            "data": {"generation": {"units": "thousand megawatthours"}},
            "frequency": [{"id": "monthly"}, {"id": "annual"}],
        },
    }
    (HERE / "eia" / "eia_v2_route_metadata_SYNTHETIC.json").write_text(json.dumps(metadata, indent=1) + "\n")
    for f in ("eia/eia_v2_electric_power_operational_data_SYNTHETIC.json", "eia/eia_v2_route_metadata_SYNTHETIC.json"):
        note(f, "SYNTHETIC", "shape of https://api.eia.gov/v2/electricity/electric-power-operational-data",
             "Values invented. Facet ids location/sectorid/fueltypeid and data id 'generation' are verified "
             "(third-party code calling the live API; datasetiq query string); the full response shape was not seen.")


# ======================================================================================
# NESO: EA register XLSX + TEC register CSV (SYNTHETIC values)
# ======================================================================================

TEC_HEADER = [  # 15 fields from the NESO data-portal data dictionary (2026-10-20)
    "Project Name", "Customer Name", "Connection Site", "Stage", "MW Connected", "MW Increase / Decrease",
    "Cumulative Total Capacity (MW)", "MW Effective From", "Project Status", "Agreement Type", "HOST TO",
    "Plant Type", "Project ID", "Project Number", "Gate",
]


def build_neso() -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "EA Register"
    ws.append(["Existing Agreement Register - SYNTHETIC FIXTURE (only columns A and F are verified)"])
    ws.append([])
    ws.append(["Project Name", "UNSEEN COLUMN B", "UNSEEN COLUMN C", "UNSEEN COLUMN D", "UNSEEN COLUMN E",
               "Technology Type", "Capacity MW (UNSEEN NAME)", "Connection Site (UNSEEN NAME)"])
    rows = [
        ("SYNTHETIC North Data Centre Campus", "Transmission Connected Demand", 300, "SYNTH SITE A"),
        ("SYNTHETIC Hyperscale Park Phase 1", "Transmission Connected Demand", 180, "SYNTH SITE B"),
        ("SYNTHETIC Datacenter Ltd Expansion", "Transmission Connected Demand", 90, "SYNTH SITE B"),
        ("SYNTHETIC Steelworks Arc Furnace", "Transmission Connected Demand", 250, "SYNTH SITE C"),
        ("SYNTHETIC Railway Traction Supply", " Transmission Connected Demand ", 40, "SYNTH SITE D"),
        ("SYNTHETIC Wind Farm", "Wind Onshore", 200, "SYNTH SITE E"),
        ("SYNTHETIC Data Centre Battery", "Energy Storage System", 50, "SYNTH SITE A"),
    ]
    for name, tech, mw, site in rows:
        ws.append([name, None, None, None, None, tech, mw, site])
    save_xlsx(wb, HERE / "neso" / "ea_register_SYNTHETIC.xlsx")
    note("neso/ea_register_SYNTHETIC.xlsx", "SYNTHETIC",
         "layout hint for https://www.neso.energy/document/373996/download (binary; not readable via WebFetch)",
         "Only 'Project Name' (column A) and 'Technology Type' (column F) with value 'Transmission Connected "
         "Demand' are verified (NESO FOI/25/174 response, 2025-12-18). All other headers are deliberately "
         "labelled UNSEEN. Values invented.")

    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(TEC_HEADER)
    w.writerow(["SYNTHETIC Wind Farm", "SYNTHETIC LTD", "SYNTH SITE E", "1", "0", "200", "200",
                "01/04/2028", "Awaiting Consents", "Directly Connected", "SYNTH TO", "Wind Onshore", "SYN-1", "1", "2"])
    w.writerow(["SYNTHETIC Data Centre Energy Centre", "SYNTHETIC DC LTD", "SYNTH SITE A", "1", "0", "60", "60",
                "01/10/2029", "Scoping", "Directly Connected", "SYNTH TO", "SYNTHETIC GAS", "SYN-2", "2", "2"])
    (HERE / "neso" / "tec_register_SYNTHETIC.csv").write_text(buf.getvalue())
    note("neso/tec_register_SYNTHETIC.csv", "SYNTHETIC",
         "layout of the TEC register CSV (https://www.neso.energy/data-portal/transmission-entry-capacity-tec-register)",
         "15 column names from the portal's data dictionary; the CSV itself is robots-blocked for WebFetch. "
         "Values invented; date and Stage formats not observed.")


def main() -> None:
    build_ercot()
    build_spp()
    build_miso()
    build_nyiso()
    build_pjm()
    build_eia()
    build_neso()
    (HERE / "provenance.json").write_text(json.dumps(PROVENANCE, indent=2, sort_keys=True) + "\n")
    print(f"Wrote {len(PROVENANCE)} fixture files under {HERE}")


if __name__ == "__main__":
    main()
